import asyncio
import base64
import io
from pathlib import Path
from uuid import UUID, uuid4
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from fastapi import UploadFile
from PIL import Image, ImageOps, UnidentifiedImageError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.config import settings
from app.models.chat import ChatAttachment

IMAGE_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}
FILE_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".json": "application/json",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class AttachmentValidationError(ValueError):
    pass


def attachment_path(storage_name: str) -> Path:
    return Path(settings.upload_dir).resolve() / storage_name


def safe_filename(filename: str | None) -> str:
    name = Path(filename or "attachment").name.strip().replace("\x00", "")
    return (name or "attachment")[:255]


def validate_kind_and_type(filename: str, requested_kind: str) -> tuple[str, str]:
    extension = Path(filename).suffix.lower()
    if requested_kind == "image" and extension in IMAGE_TYPES:
        return extension, IMAGE_TYPES[extension]
    if requested_kind == "file" and extension in FILE_TYPES:
        return extension, FILE_TYPES[extension]
    if requested_kind == "image":
        raise AttachmentValidationError("이미지는 JPG, PNG, WebP 형식만 첨부할 수 있습니다.")
    raise AttachmentValidationError(
        "파일은 TXT, MD, CSV, JSON, PDF, DOCX 형식만 첨부할 수 있습니다."
    )


def verify_image(data: bytes) -> None:
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
    except (UnidentifiedImageError, OSError, ValueError) as error:
        raise AttachmentValidationError("올바른 이미지 파일이 아닙니다.") from error


def verify_document(data: bytes, content_type: str) -> None:
    try:
        if content_type == "application/pdf":
            PdfReader(io.BytesIO(data))
        elif content_type.endswith("wordprocessingml.document"):
            Document(io.BytesIO(data))
    except (BadZipFile, KeyError, OSError, PackageNotFoundError, PdfReadError, ValueError) as error:
        raise AttachmentValidationError("손상되었거나 지원하지 않는 문서 파일입니다.") from error


async def save_upload(upload: UploadFile, requested_kind: str, user_id: UUID) -> ChatAttachment:
    filename = safe_filename(upload.filename)
    extension, content_type = validate_kind_and_type(filename, requested_kind)
    data = await upload.read(settings.upload_max_bytes + 1)
    await upload.close()
    if not data:
        raise AttachmentValidationError("빈 파일은 첨부할 수 없습니다.")
    if len(data) > settings.upload_max_bytes:
        limit_mb = settings.upload_max_bytes // (1024 * 1024)
        raise AttachmentValidationError(f"파일 하나의 최대 크기는 {limit_mb}MB입니다.")
    if requested_kind == "image":
        await asyncio.to_thread(verify_image, data)
    elif content_type in {
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }:
        await asyncio.to_thread(verify_document, data, content_type)

    upload_dir = Path(settings.upload_dir).resolve()
    await asyncio.to_thread(upload_dir.mkdir, parents=True, exist_ok=True)
    storage_name = f"{user_id}/{uuid4().hex}{extension}"
    destination = upload_dir / storage_name
    await asyncio.to_thread(destination.parent.mkdir, parents=True, exist_ok=True)
    await asyncio.to_thread(destination.write_bytes, data)
    return ChatAttachment(
        user_id=user_id,
        kind=requested_kind,
        original_filename=filename,
        storage_name=storage_name,
        content_type=content_type,
        size_bytes=len(data),
    )


async def remove_stored_file(attachment: ChatAttachment) -> None:
    path = attachment_path(attachment.storage_name)
    await asyncio.to_thread(path.unlink, missing_ok=True)


def _extract_text(path: Path, content_type: str) -> str:
    if content_type.startswith("text/") or content_type == "application/json":
        return path.read_text(encoding="utf-8", errors="replace")
    if content_type == "application/pdf":
        reader = PdfReader(path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if content_type.endswith("wordprocessingml.document"):
        document = Document(path)
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    return ""


async def extracted_attachment_text(attachment: ChatAttachment) -> str:
    try:
        text = await asyncio.to_thread(
            _extract_text,
            attachment_path(attachment.storage_name),
            attachment.content_type,
        )
    except (BadZipFile, KeyError, OSError, PackageNotFoundError, PdfReadError, ValueError):
        return ""
    return text.strip()[: settings.attachment_text_max_characters]


def _vlm_image_payload(path: Path, content_type: str) -> tuple[bytes, str]:
    data = path.read_bytes()
    with Image.open(io.BytesIO(data)) as source:
        source.load()
        image = ImageOps.exif_transpose(source)
        needs_resize = max(image.size) > settings.vlm_image_max_dimension
        needs_compression = len(data) > settings.vlm_image_max_bytes
        if not needs_resize and not needs_compression:
            return data, content_type

        image.thumbnail(
            (settings.vlm_image_max_dimension, settings.vlm_image_max_dimension),
            Image.Resampling.LANCZOS,
        )
        if image.mode in {"RGBA", "LA"} or (
            image.mode == "P" and "transparency" in image.info
        ):
            rgba = image.convert("RGBA")
            background = Image.new("RGBA", rgba.size, "white")
            background.alpha_composite(rgba)
            image = background.convert("RGB")
        else:
            image = image.convert("RGB")

        output = io.BytesIO()
        image.save(
            output,
            format="JPEG",
            quality=settings.vlm_image_jpeg_quality,
            optimize=True,
        )
        return output.getvalue(), "image/jpeg"


async def image_data_url(attachment: ChatAttachment) -> str:
    data, content_type = await asyncio.to_thread(
        _vlm_image_payload,
        attachment_path(attachment.storage_name),
        attachment.content_type,
    )
    encoded = base64.b64encode(data).decode("ascii")
    return f"data:{content_type};base64,{encoded}"
